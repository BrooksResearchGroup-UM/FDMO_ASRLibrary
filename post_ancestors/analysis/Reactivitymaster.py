#!/usr/bin/env python
# -*- coding: utf-8 -*-

__author__ = "Aidan Cosgrove <aidancos@umich.edu>"

"""
ReactivityMaster.py

1. Consolidates all substrate reactivity spreadsheets into one master DataFrame
2. Produces per-substrate per-plate heatmaps using the existing color scheme
3. Saves master CSV for latent space overlay

Sheet naming conventions:
  - Files with 'G3' in name: sheets are gen 3 pt1, pt2, pt3
  - All other files: sheets are gen 1, gen 2, gen 3 pt1, pt2, pt3

Color scheme: white → #115472 (opacity gradient, all ancestors for now)
"""

import os
import re
import numpy as np
import pandas as pd
import matplotlib as mpl
mpl.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.colors as mcolors
import openpyxl

################################################
# Paths
################################################

DATA_DIR    = os.path.expanduser("~/KJ_G3Paper/KJ_G3PAPER/plate_heatmaps")
RESULTS_DIR = os.path.join(os.path.dirname(DATA_DIR), "heatmap_output")
MASTER_CSV  = os.path.join(os.path.dirname(DATA_DIR), "master_reactivity.csv")
os.makedirs(RESULTS_DIR, exist_ok=True)

################################################
# Color palette
################################################

ANCESTOR_CMAP = mcolors.LinearSegmentedColormap.from_list(
    "ancestor", [(1, 1, 1, 1), mcolors.to_rgba("#115472", 1)]
)
ROW_LABELS = list("ABCDEFGH")
COL_LABELS = [str(i) for i in range(1, 13)]

################################################
# Sheet name normalizer
################################################

def normalize_sheet_name(sheet_name, is_g3_file):
    # strip common prefixes before matching
    name = sheet_name.lower().strip()
    name = re.sub(r'^para_|^ph_|^2ph_|^3me_', '', name)

    if re.search(r'g3.*p3|gen.*3.*p.*3|3.*p.*3', name):
        return "gen3p3"
    if re.search(r'g3.*p2|gen.*3.*p.*2|3.*p.*2', name):
        return "gen3p2"
    if re.search(r'g3.*p1|gen.*3.*p.*1|3.*p.*1', name):
        return "gen3p1"
    if is_g3_file:
        return None
    if re.search(r'\bg1\b|gen.*1|generation.*1', name):
        return "gen1"
    if re.search(r'\bg2\b|gen.*2|generation.*2', name):
        return "gen2"
    return None

def get_generation_labels(sheet_names, is_g3_file):
    """
    Map sheet names to generation labels in order.
    Falls back to positional mapping if regex fails.
    """
    if is_g3_file:
        positional = ["gen3p1", "gen3p2", "gen3p3"]
    else:
        positional = ["gen1", "gen2", "gen3p1", "gen3p2", "gen3p3"]

    labels = []
    for i, name in enumerate(sheet_names):
        label = normalize_sheet_name(name, is_g3_file)
        if label is None and i < len(positional):
            label = positional[i]
        labels.append(label)
    return labels

################################################
# Plate parser (same logic as plate_heatmaps.py)
################################################

def parse_plate_positional(ws):
    """
    For sheets with no row labels — finds the 8×12 numeric grid positionally.
    """
    plate = np.zeros((8, 12))
    numeric_rows = []
    for row in ws.iter_rows(values_only=True):
        nums = [v for v in row if isinstance(v, (int, float))]
        if len(nums) >= 12:
            numeric_rows.append(nums[:12])
    for i, nums in enumerate(numeric_rows[:8]):
        plate[i] = nums
    return plate

def parse_plate(ws):
    plate     = np.zeros((8, 12))
    row_order = {r: i for i, r in enumerate(ROW_LABELS)}

    for row in ws.iter_rows(values_only=True):
        label_col = None
        row_label = None
        for ci, val in enumerate(row):
            if val in row_order:
                label_col = ci
                row_label = val
                break
            # handle 'gen1_A01' style
            if isinstance(val, str) and '_' in val:
                parts = val.split('_')
                if len(parts) >= 2 and len(parts[-1]) >= 1:
                    candidate = parts[-1][0].upper()
                    if candidate in row_order:
                        label_col = ci
                        row_label = candidate
                        break
        if label_col is None:
            continue

        row_idx = row_order[row_label]
        vals = []
        for val in row[label_col + 1:]:
            if isinstance(val, (int, float)):
                vals.append(float(val))
            if len(vals) == 12:
                break
        if len(vals) == 12:
            plate[row_idx] = vals

    return plate

def get_plate_data(wb, fname):
    sheet_names = [s for s in wb.sheetnames
                   if 'template' not in s.lower()]

    all_sheet = next((s for s in sheet_names
                      if s.lower().strip() in ('all', 'all gen 3', 'all gens', 'all gen3')), None)

    if all_sheet:
        ws    = wb[all_sheet]
        fmt   = detect_format(ws)
        plate = parse_plate(ws) if fmt == 'labeled' else \
                parse_plate_wellid(ws) if fmt == 'wellid' else \
                parse_plate_positional(ws)
        if plate.max() == 0 and fmt == 'labeled':
            plate = parse_plate_positional(ws)
        print(f"  Using '{all_sheet}' sheet  (max={plate.max():.4f})")
        return {"all": plate}

    is_g3_file  = 'G3' in fname or 'g3' in fname
    data_sheets = [s for s in sheet_names
                   if s.lower().strip() not in ('all', 'all gen 3', 'all gens', 'all gen3')
                   and 'raw' not in s.lower()]
    gen_labels  = get_generation_labels(data_sheets, is_g3_file)

    result = {}
    for sheet_name, gen_label in zip(data_sheets, gen_labels):
        if gen_label is None:
            print(f"  ⚠ Could not map sheet '{sheet_name}' — skipping")
            continue
        ws  = wb[sheet_name]
        fmt = detect_format(ws)
        plate = parse_plate(ws) if fmt == 'labeled' else \
                parse_plate_wellid(ws) if fmt == 'wellid' else \
                parse_plate_positional(ws)
        if plate.max() == 0 and fmt == 'labeled':
            plate = parse_plate_positional(ws)
        print(f"  {sheet_name} → {gen_label}  (max={plate.max():.4f})")
        result[gen_label] = plate

    return result


def main():
    xlsx_files = [f for f in os.listdir(DATA_DIR)
                  if f.endswith('.xlsx') and not f.startswith('~')]

    print(f"Found {len(xlsx_files)} xlsx files in {DATA_DIR}")

    all_plates = {}
    global_max = 0

    for fname in sorted(xlsx_files):
        substrate = fname.replace('.xlsx', '').replace('_', ' ').replace('-', ' ')
        fpath = os.path.join(DATA_DIR, fname)
        try:
            wb = openpyxl.load_workbook(fpath, data_only=True)
        except Exception as e:
            print(f"  ⚠ Could not open {fname}: {e}")
            continue

        print(f"\n{fname}")
        plates = get_plate_data(wb, fname)

        for gen_label, plate in plates.items():
            key = (substrate, gen_label)
            all_plates[key] = plate
            if plate.max() > 0:
                global_max = max(global_max,
                                 np.percentile(plate[plate > 0], 95))

    print(f"\nGlobal vmax (95th percentile): {global_max:.4f}")

    for (substrate, gen_label), plate in all_plates.items():
        safe_sub  = substrate.replace(' ', '_').replace('/', '_')
        out_fname = f"{safe_sub}_{gen_label}_heatmap.png"
        out_path  = os.path.join(RESULTS_DIR, out_fname)
        title     = f"{substrate} — {gen_label.replace('gen', 'Gen ').replace('p', ' Pt.')}"
        plot_plate(plate, title, ANCESTOR_CMAP, out_path,
                   vmin=0, vmax=global_max)
        print(f"  Saved: {out_fname}")

    records = []
    for (substrate, gen_label), plate in all_plates.items():
        for r, row_label in enumerate(ROW_LABELS):
            for c, col_label in enumerate(COL_LABELS):
                records.append({
                    'substrate':  substrate,
                    'generation': gen_label,
                    'well':       f"{row_label}{col_label}",
                    'row':        row_label,
                    'col':        int(col_label),
                    'plate_row':  r,
                    'plate_col':  c,
                    'reactivity': plate[r, c],
                    'active':     plate[r, c] > 0,
                })

    master_df = pd.DataFrame(records)
    master_df.to_csv(MASTER_CSV, index=False)
    print(f"\nMaster CSV saved: {MASTER_CSV}")
    print(f"Total wells: {len(master_df)}")
    print(f"Active wells: {master_df['active'].sum()}")
    print(f"\nSubstrates found:")
    for sub in sorted(master_df['substrate'].unique()):
        n_active = master_df[master_df['substrate'] == sub]['active'].sum()
        print(f"  {sub}: {n_active} active wells")

    return master_df

def parse_plate_wellid(ws, value_col=None):
    """
    For sheets where rows are individual wells (A01, A02... H12)
    and reactivity is in a specific column.
    value_col: if None, uses the last numeric column.
    """
    plate = np.zeros((8, 12))
    row_order = {r: i for i, r in enumerate(ROW_LABELS)}

    for row in ws.iter_rows(values_only=True):
        well_id = row[0]
        if not isinstance(well_id, str) or len(well_id) < 3:
            continue
        row_letter = well_id[0].upper()
        if row_letter not in row_order:
            continue
        try:
            col_num = int(well_id[1:]) - 1  # A01 → col 0
        except ValueError:
            continue
        if col_num < 0 or col_num > 11:
            continue

        # get reactivity value
        nums = [(ci, v) for ci, v in enumerate(row[1:], 1)
                if isinstance(v, (int, float))]
        if not nums:
            continue
        if value_col is not None:
            val = row[value_col] if isinstance(row[value_col], (int, float)) else 0
        else:
            # use last numeric column — prod/IS or minus background
            val = nums[-1][1]

        plate[row_order[row_letter], col_num] = max(0, val)

    return plate

def detect_format(ws):
    """Returns 'wellid', 'labeled', or 'positional'"""
    for row in ws.iter_rows(values_only=True):
        for val in row:
            if isinstance(val, str):
                if len(val) >= 3 and val[0].upper() in 'ABCDEFGH':
                    try:
                        int(val[1:3])
                        return 'wellid'
                    except ValueError:
                        pass
                if val in list('ABCDEFGH'):
                    return 'labeled'
    return 'positional'
################################################
# Heatmap plotter
################################################

def plot_plate(plate, title, cmap, out_path, vmin=0, vmax=1.0):
    fig, ax = plt.subplots(figsize=(10, 5))
    fig.patch.set_facecolor("white")

    im = ax.imshow(plate, cmap=cmap, vmin=vmin, vmax=vmax, aspect='auto')

    for x in np.arange(-0.5, 12, 1):
        ax.axvline(x, color='white', lw=0.8)
    for y in np.arange(-0.5, 8, 1):
        ax.axhline(y, color='white', lw=0.8)

    for r in range(8):
        for c in range(12):
            val = plate[r, c]
            text_color = 'white' if val > vmax * 0.55 else '#444444'
            label = f"{val:.3f}" if val > 0 else ""
            ax.text(c, r, label, ha='center', va='center',
                    fontsize=5.5, color=text_color, fontweight='500')

    ax.set_xticks(range(12))
    ax.set_xticklabels(COL_LABELS, fontsize=9)
    ax.set_yticks(range(8))
    ax.set_yticklabels(ROW_LABELS, fontsize=9)
    ax.xaxis.tick_top()
    ax.xaxis.set_label_position('top')

    ax.set_title(title, fontsize=12, fontweight='bold', pad=14, color='#1a1a1a')

    cbar = plt.colorbar(im, ax=ax, fraction=0.025, pad=0.02)
    cbar.set_label("Reactivity (prod/IS)", fontsize=8, color='#444')
    cbar.ax.tick_params(labelsize=7)

    plt.tight_layout()
    plt.savefig(out_path, dpi=200, bbox_inches='tight',
                facecolor='white', edgecolor='none')
    plt.close()

################################################
# Main consolidation pipeline
################################################

def main():
    xlsx_files = [f for f in os.listdir(DATA_DIR)
                  if f.endswith('.xlsx') and not f.startswith('~')]

    print(f"Found {len(xlsx_files)} xlsx files in {DATA_DIR}")

    all_plates = {}
    global_max = 0

    for fname in sorted(xlsx_files):
        substrate = fname.replace('.xlsx', '').replace('_', ' ').replace('-', ' ')
        fpath = os.path.join(DATA_DIR, fname)
        try:
            wb = openpyxl.load_workbook(fpath, data_only=True)
        except Exception as e:
            print(f"  ⚠ Could not open {fname}: {e}")
            continue

        print(f"\n{fname}")
        plates = get_plate_data(wb, fname)

        for gen_label, plate in plates.items():
            key = (substrate, gen_label)
            all_plates[key] = plate
            if plate.max() > 0:
                global_max = max(global_max,
                                 np.percentile(plate[plate > 0], 95))

    print(f"\nGlobal vmax (95th percentile): {global_max:.4f}")

    for (substrate, gen_label), plate in all_plates.items():
        safe_sub  = substrate.replace(' ', '_').replace('/', '_')
        out_fname = f"{safe_sub}_{gen_label}_heatmap.png"
        out_path  = os.path.join(RESULTS_DIR, out_fname)
        title     = f"{substrate} — {gen_label.replace('gen', 'Gen ').replace('p', ' Pt.')}"
        plot_plate(plate, title, ANCESTOR_CMAP, out_path,
                   vmin=0, vmax=global_max)
        print(f"  Saved: {out_fname}")

    records = []
    for (substrate, gen_label), plate in all_plates.items():
        for r, row_label in enumerate(ROW_LABELS):
            for c, col_label in enumerate(COL_LABELS):
                records.append({
                    'substrate':  substrate,
                    'generation': gen_label,
                    'well':       f"{row_label}{col_label}",
                    'row':        row_label,
                    'col':        int(col_label),
                    'plate_row':  r,
                    'plate_col':  c,
                    'reactivity': plate[r, c],
                    'active':     plate[r, c] > 0,
                })

    master_df = pd.DataFrame(records)
    master_df.to_csv(MASTER_CSV, index=False)
    print(f"\nMaster CSV saved: {MASTER_CSV}")
    print(f"Total wells: {len(master_df)}")
    print(f"Active wells: {master_df['active'].sum()}")
    print(f"\nSubstrates found:")
    for sub in sorted(master_df['substrate'].unique()):
        n_active = master_df[master_df['substrate'] == sub]['active'].sum()
        print(f"  {sub}: {n_active} active wells")

    return master_df


if __name__ == "__main__":
    main()